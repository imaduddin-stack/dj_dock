#!/usr/bin/env python
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

def add_dokumen_lkps_sheet():
    """Add Dokumen-LKPS sheet template to seed_template.xlsx"""

    # Load workbook
    wb = openpyxl.load_workbook('seed_template.xlsx')

    # Remove existing sheet if any
    if 'Dokumen-LKPS' in wb.sheetnames:
        del wb['Dokumen-LKPS']

    # Create new sheet
    ws = wb.create_sheet('Dokumen-LKPS')

    # Define styles
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)
    border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )

    # Write headers
    headers = ['dokumen_id', 'lkps_id', 'keterangan']
    for col_num, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_num)
        cell.value = header
        cell.fill = header_fill
        cell.font = header_font
        cell.border = border
        cell.alignment = Alignment(horizontal='center', vertical='center')

    # Write sample data (empty rows with borders)
    for row_num in range(2, 12):  # 10 empty rows for data entry
        for col_num in range(1, 4):
            cell = ws.cell(row=row_num, column=col_num)
            cell.border = border
            cell.alignment = Alignment(horizontal='left', vertical='top')

    # Set column widths
    ws.column_dimensions['A'].width = 15
    ws.column_dimensions['B'].width = 15
    ws.column_dimensions['C'].width = 50

    # Add instruction comment
    instruction = """FORMAT PENGISIAN:

kolom A (dokumen_id): ID dari tabel Dokumen
kolom B (lkps_id): ID dari tabel LKPS Elemen
kolom C (keterangan): Catatan tambahan (opsional)

CONTOH:
dokumen_id | lkps_id | keterangan
-----------|---------|----------------------------
1          | 1       | Dokumen pendukung Tabel 1.a
1          | 2       | Dokumen pendukung Tabel 1.b
2          | 5       | Laporan implementasi
"""

    ws['E1'].value = instruction
    ws['E1'].alignment = Alignment(horizontal='left', vertical='top', wrap_text=True)

    # Save
    wb.save('seed_template.xlsx')
    print("[OK] Sheet 'Dokumen-LKPS' berhasil ditambahkan ke seed_template.xlsx")
    print("\nPanduan:")
    print("  1. Buka file seed_template.xlsx")
    print("  2. Buka sheet 'Dokumen-LKPS'")
    print("  3. Isi relasi dokumen-lkps di kolom A dan B")
    print("  4. Jalankan: python manage.py import_dokumen_lkps seed_template.xlsx")

if __name__ == '__main__':
    add_dokumen_lkps_sheet()

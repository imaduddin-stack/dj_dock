import json
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from pathlib import Path

# Load seed data
with open('seed_data.json', 'r', encoding='utf-8') as f:
    data = json.load(f)

# Create workbook
wb = openpyxl.Workbook()
wb.remove(wb.active)  # Remove default sheet

# Define header style
header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
header_font = Font(bold=True, color="FFFFFF")

def create_sheet(sheet_name, data_list, columns):
    """Create a sheet with headers and data"""
    ws = wb.create_sheet(sheet_name)

    # Add header row
    for col_idx, column in enumerate(columns, 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.value = column
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    # Add data rows
    for row_idx, item in enumerate(data_list, 2):
        for col_idx, column in enumerate(columns, 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.value = item.get(column, '')
            cell.alignment = Alignment(horizontal="left", vertical="top")

    # Auto adjust column width
    for col_idx, column in enumerate(columns, 1):
        max_length = max(
            len(str(column)),
            max([len(str(item.get(column, ''))) for item in data_list] or [0])
        )
        ws.column_dimensions[openpyxl.utils.get_column_letter(col_idx)].width = min(max_length + 2, 50)

# 1. Create Kriteria sheet
create_sheet(
    'Kriteria',
    data['kriteria'],
    ['id', 'nomor', 'nama']
)

# 2. Create LED Komponen sheet
create_sheet(
    'LED Komponen',
    data['led_komponen'],
    ['id', 'kriteria_id', 'kode', 'deskripsi']
)

# 3. Create LKPS Elemen sheet
create_sheet(
    'LKPS Elemen',
    data['lkps_elemen'],
    ['id', 'kriteria_id', 'tabel_referensi', 'deskripsi']
)

# 4. Create Dokumen sheet
dokumen_data = []
for doc in data['dokumen']:
    doc_copy = doc.copy()
    # Convert lists to comma-separated strings for Excel
    doc_copy['dipakai_led'] = ', '.join(map(str, doc.get('dipakai_led', [])))
    doc_copy['dipakai_lkps'] = ', '.join(map(str, doc.get('dipakai_lkps', [])))
    dokumen_data.append(doc_copy)

create_sheet(
    'Dokumen',
    dokumen_data,
    ['id', 'judul', 'jenis', 'deskripsi', 'link_gdrive', 'dipakai_led', 'dipakai_lkps']
)

# Save file
output_path = 'seed_template.xlsx'
wb.save(output_path)
print(f"[OK] File Excel berhasil dibuat: {output_path}")
print(f"     Sheets yang dibuat:")
print(f"     - Kriteria")
print(f"     - LED Komponen")
print(f"     - LKPS Elemen")
print(f"     - Dokumen")

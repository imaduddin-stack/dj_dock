#!/usr/bin/env python
import os
import sys
import django
import openpyxl
from pathlib import Path

# Setup Django
os.environ.setdefault('DJANGO_SETTINGS_MODULE', 'pste.settings')
sys.path.insert(0, os.path.join(os.path.dirname(__file__), 'pste'))
django.setup()

from dokumen.models import LkpsElemen, Kriteria

def import_lkps_elemen_from_excel(excel_file, sheet_name='LKPS_Elemen'):
    """Import LKPS Elemen data dari Excel"""

    if not os.path.exists(excel_file):
        print(f"ERROR: File {excel_file} tidak ditemukan!")
        return False

    wb = openpyxl.load_workbook(excel_file)

    if sheet_name not in wb.sheetnames:
        print(f"ERROR: Sheet '{sheet_name}' tidak ditemukan!")
        print(f"Available sheets: {', '.join(wb.sheetnames)}")
        return False

    ws = wb[sheet_name]

    success_count = 0
    skip_count = 0
    error_count = 0
    errors_log = []

    print(f"\n{'='*70}")
    print(f"IMPORT LKPS ELEMEN FROM EXCEL")
    print(f"File: {excel_file}")
    print(f"Sheet: {sheet_name}")
    print(f"{'='*70}\n")

    # Get headers (row 1)
    headers = []
    for cell in ws[1]:
        headers.append(cell.value)

    print(f"Headers: {headers}\n")

    # Find column indices
    try:
        id_col = headers.index('ID') + 1
        kriteria_no_col = headers.index('Kriteria No') + 1
        tabel_ref_col = headers.index('Tabel Ref') + 1
        deskripsi_col = headers.index('Deskripsi') + 1
    except ValueError as e:
        print(f"ERROR: Column not found - {str(e)}")
        return False

    print(f"Column mapping:")
    print(f"  ID (Col {id_col})")
    print(f"  Kriteria No (Col {kriteria_no_col})")
    print(f"  Tabel Ref (Col {tabel_ref_col})")
    print(f"  Deskripsi (Col {deskripsi_col})\n")

    # Iterate rows (skip header)
    for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=False), start=2):
        try:
            # Get values
            lkps_id = row[id_col - 1].value
            kriteria_nomor = row[kriteria_no_col - 1].value
            tabel_referensi = row[tabel_ref_col - 1].value
            deskripsi = row[deskripsi_col - 1].value

            # Skip if empty row
            if not lkps_id or not tabel_referensi:
                skip_count += 1
                continue

            # Find Kriteria by nomor
            try:
                kriteria = Kriteria.objects.get(nomor=kriteria_nomor)
            except Kriteria.DoesNotExist:
                # Skip invalid kriteria
                skip_count += 1
                errors_log.append(f"Row {row_idx}: Kriteria '{kriteria_nomor}' tidak ditemukan (SKIP)")
                print(f"  [SKIP] Row {row_idx}: Kriteria '{kriteria_nomor}' tidak ditemukan")
                continue

            # Create or update LkpsElemen
            lkps_obj, created = LkpsElemen.objects.update_or_create(
                id=lkps_id,
                defaults={
                    'kriteria': kriteria,
                    'tabel_referensi': tabel_referensi,
                    'deskripsi': deskripsi or ''
                }
            )

            if created:
                success_count += 1
                print(f"  [CREATE] Row {row_idx}: ID {lkps_id} - {tabel_referensi}")
            else:
                success_count += 1
                print(f"  [UPDATE] Row {row_idx}: ID {lkps_id} - {tabel_referensi}")

        except Exception as e:
            error_count += 1
            errors_log.append(f"Row {row_idx}: {str(e)}")
            print(f"  [ERROR] Row {row_idx}: {str(e)}")
            continue

    # Print summary
    print(f"\n{'='*70}")
    print(f"SUMMARY:")
    print(f"  Created/Updated: {success_count}")
    print(f"  Skipped: {skip_count}")
    print(f"  Errors: {error_count}")
    print(f"{'='*70}\n")

    if errors_log:
        print(f"ERROR LOG:")
        for error in errors_log:
            print(f"  - {error}")
        print()

    # Verify final count
    final_count = LkpsElemen.objects.count()
    print(f"Final LKPS Elemen count in database: {final_count}")
    print()

    return error_count == 0


def main():
    excel_file = "database_dump_all_20260626_103022.xlsx"
    sheet_name = "3_LKPS_Elemen"

    success = import_lkps_elemen_from_excel(excel_file, sheet_name)

    if success:
        print("[OK] Import completed successfully!")
        sys.exit(0)
    else:
        print("[ERROR] Import completed with errors!")
        sys.exit(1)


if __name__ == '__main__':
    main()

import openpyxl
from django.core.management.base import BaseCommand, CommandError
from dokumen.models import Dokumen, LkpsElemen, DokumenLkps


class Command(BaseCommand):
    help = 'Import relasi Dokumen-LKPS dari Excel file (sheet: Dokumen-LKPS)'

    def add_arguments(self, parser):
        parser.add_argument(
            'excel_file',
            type=str,
            help='Path ke file Excel yang berisi sheet "Dokumen-LKPS"'
        )

    def handle(self, *args, **options):
        excel_file = options['excel_file']

        try:
            wb = openpyxl.load_workbook(excel_file)
        except FileNotFoundError:
            raise CommandError(f'File tidak ditemukan: {excel_file}')
        except Exception as e:
            raise CommandError(f'Error membaca file: {str(e)}')

        # Check if sheet exists
        if 'Dokumen-LKPS' not in wb.sheetnames:
            raise CommandError(
                f'Sheet "Dokumen-LKPS" tidak ditemukan. '
                f'Sheet yang tersedia: {", ".join(wb.sheetnames)}'
            )

        ws = wb['Dokumen-LKPS']
        success_count = 0
        skip_count = 0
        error_count = 0

        self.stdout.write(self.style.SUCCESS(f'\n=== Import Relasi Dokumen-LKPS ==='))
        self.stdout.write(f'File: {excel_file}\n')

        # Iterate rows (skip header)
        for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            try:
                dokumen_id = row[0]
                lkps_id = row[1]
                keterangan = row[2] if len(row) > 2 else ''

                # Skip jika dokumen_id atau lkps_id tidak ada
                if not dokumen_id or not lkps_id:
                    skip_count += 1
                    continue

                # Cari dokumen dan lkps_elemen
                try:
                    dokumen = Dokumen.objects.get(id=dokumen_id)
                except Dokumen.DoesNotExist:
                    error_count += 1
                    self.stdout.write(
                        self.style.ERROR(f'Row {row_idx}: Dokumen ID {dokumen_id} tidak ditemukan')
                    )
                    continue

                try:
                    lkps_elemen = LkpsElemen.objects.get(id=lkps_id)
                except LkpsElemen.DoesNotExist:
                    error_count += 1
                    self.stdout.write(
                        self.style.ERROR(f'Row {row_idx}: LKPS Elemen ID {lkps_id} tidak ditemukan')
                    )
                    continue

                # Create DokumenLkps
                dok_lkps, created = DokumenLkps.objects.get_or_create(
                    dokumen=dokumen,
                    lkps_elemen=lkps_elemen,
                    defaults={'keterangan': keterangan or ''}
                )

                if created:
                    success_count += 1
                    self.stdout.write(
                        self.style.SUCCESS(
                            f'✓ Row {row_idx}: {dokumen.judul} → {lkps_elemen.tabel_referensi}'
                        )
                    )

            except Exception as e:
                error_count += 1
                self.stdout.write(self.style.ERROR(f'Row {row_idx}: {str(e)}'))

        # Print summary
        self.stdout.write('\n' + '=' * 70)
        self.stdout.write(self.style.SUCCESS(f'Berhasil dibuat: {success_count}'))
        self.stdout.write(f'Skipped: {skip_count}')
        self.stdout.write(self.style.ERROR(f'Error: {error_count}'))
        self.stdout.write('=' * 70 + '\n')

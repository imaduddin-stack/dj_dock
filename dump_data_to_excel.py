#!/usr/bin/env python
import os
import sys
import django
from datetime import datetime
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

# Setup Django
os.environ.setdefault('DJANGO_SETTINGS_MODULE', 'pste.settings')
sys.path.insert(0, os.path.join(os.path.dirname(__file__), 'pste'))
django.setup()

from dokumen.models import Kriteria, LedKomponen, LkpsElemen, Dokumen, DokumenLkps

def create_excel_sheet(wb, sheet_name, headers, data_rows):
    """Helper untuk membuat sheet dengan style"""
    ws = wb.create_sheet(sheet_name)

    # Header style
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)
    border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )

    # Write headers
    for col_num, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_num)
        cell.value = header
        cell.fill = header_fill
        cell.font = header_font
        cell.border = border
        cell.alignment = Alignment(horizontal='center', vertical='center')

    # Write data
    for row_num, row_data in enumerate(data_rows, 2):
        for col_num, value in enumerate(row_data, 1):
            cell = ws.cell(row=row_num, column=col_num)
            cell.value = value
            cell.border = border
            cell.alignment = Alignment(horizontal='left', vertical='top', wrap_text=True)

    # Auto adjust column width
    for col_num, header in enumerate(headers, 1):
        max_length = len(str(header))
        column = openpyxl.utils.get_column_letter(col_num)
        for row_num in range(2, len(data_rows) + 2):
            cell = ws[f'{column}{row_num}']
            if cell.value:
                max_length = max(max_length, len(str(cell.value)))
        adjusted_width = min(max_length + 2, 50)
        ws.column_dimensions[column].width = adjusted_width

    return ws

def main():
    wb = openpyxl.Workbook()
    wb.remove(wb.active)  # Remove default sheet

    # 1. KRITERIA
    print("Exporting Kriteria...")
    kriteria_list = Kriteria.objects.all()[:10]
    kriteria_data = []
    for k in kriteria_list:
        kriteria_data.append([k.id, k.nomor, k.nama])
    create_excel_sheet(wb, "Kriteria", ["ID", "Nomor", "Nama"], kriteria_data)

    # 2. LED KOMPONEN
    print("Exporting LED Komponen...")
    komponen_list = LedKomponen.objects.select_related('kriteria').all()[:10]
    komponen_data = []
    for k in komponen_list:
        komponen_data.append([
            k.id,
            k.kriteria.nomor,
            k.kriteria.nama,
            k.kode,
            k.deskripsi
        ])
    create_excel_sheet(wb, "LED Komponen", ["ID", "Kriteria Nomor", "Kriteria Nama", "Kode", "Deskripsi"], komponen_data)

    # 3. LKPS ELEMEN
    print("Exporting LKPS Elemen...")
    elemen_list = LkpsElemen.objects.select_related('kriteria').all()[:10]
    elemen_data = []
    for e in elemen_list:
        elemen_data.append([
            e.id,
            e.kriteria.nomor,
            e.kriteria.nama,
            e.tabel_referensi,
            e.deskripsi
        ])
    create_excel_sheet(wb, "LKPS Elemen", ["ID", "Kriteria Nomor", "Kriteria Nama", "Tabel Referensi", "Deskripsi"], elemen_data)

    # 4. DOKUMEN
    print("Exporting Dokumen...")
    dokumen_list = Dokumen.objects.all()[:10]
    dokumen_data = []
    for d in dokumen_list:
        led_list = ", ".join([str(l) for l in d.dipakai_led.all()])
        dokumen_data.append([
            d.id,
            d.judul,
            d.get_jenis_display(),
            d.deskripsi[:100] if d.deskripsi else "",  # First 100 chars
            d.link_gdrive,
            led_list,
            d.created_at.strftime("%Y-%m-%d %H:%M:%S") if d.created_at else ""
        ])
    create_excel_sheet(wb, "Dokumen",
        ["ID", "Judul", "Jenis", "Deskripsi (100 char)", "Link GDrive", "Dipakai LED", "Created At"],
        dokumen_data)

    # 5. DOKUMEN - LKPS ELEMEN (Relasi)
    print("Exporting Dokumen-LKPS Elemen...")
    dokumen_lkps_list = DokumenLkps.objects.select_related('dokumen', 'lkps_elemen', 'lkps_elemen__kriteria').all()[:10]
    dokumen_lkps_data = []
    for dl in dokumen_lkps_list:
        dokumen_lkps_data.append([
            dl.id,
            dl.dokumen.judul,
            dl.lkps_elemen.tabel_referensi,
            dl.lkps_elemen.kriteria.nomor,
            dl.lkps_elemen.kriteria.nama,
            dl.keterangan,
            dl.created_at.strftime("%Y-%m-%d %H:%M:%S") if dl.created_at else ""
        ])
    create_excel_sheet(wb, "Dokumen-LKPS",
        ["ID", "Dokumen", "LKPS Tabel Ref", "Kriteria No", "Kriteria Nama", "Keterangan", "Created At"],
        dokumen_lkps_data)

    # Save file
    from datetime import datetime
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    output_file = f"dump_{timestamp}.xlsx"
    wb.save(output_file)
    print(f"\n[OK] Data berhasil diexport ke: {output_file}")

    # Summary
    print(f"\nRingkasan:")
    print(f"  - Kriteria: {len(kriteria_data)} records")
    print(f"  - LED Komponen: {len(komponen_data)} records")
    print(f"  - LKPS Elemen: {len(elemen_data)} records")
    print(f"  - Dokumen: {len(dokumen_data)} records")
    print(f"  - Dokumen-LKPS Elemen: {len(dokumen_lkps_data)} records")

if __name__ == '__main__':
    main()
